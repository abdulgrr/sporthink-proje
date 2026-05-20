# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_report():
    template_path = r"c:\Users\zarp\Desktop\bitirme project\proje raporu\YBS4002 Bitirme Proje Rapor Şablonu-2026.docx"
    output_path = r"c:\Users\zarp\Desktop\bitirme project\proje raporu\YBS4002_Bitirme_Proje_Raporu_Guncel.docx"
    img_dir = r"c:\Users\zarp\Desktop\bitirme project\proje raporu"

    print("Şablon yükleniyor...")
    doc = Document(template_path)

    # 1. Başlık sayfasını koruyup geri kalan geçici yönergeleri temizleyelim.
    print("Yönergeler temizleniyor...")
    total_paragraphs = len(doc.paragraphs)
    for _ in range(total_paragraphs - 31):
        p = doc.paragraphs[31]
        p._element.getparent().remove(p._element)

    # Tabloları temizleyelim
    for t in list(doc.tables):
        t._element.getparent().remove(t._element)

    # Ortak paragraf formatlama yardımcı fonksiyonu
    def add_p(text, style='Normal', bold=False, italic=False, space_after=6, line_spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.bold = bold
        run.italic = italic
        return p

    def add_heading(text, level, space_before=12, space_after=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.bold = True
        
        if level == 1:
            run.font.size = Pt(14)
            p.paragraph_format.space_before = Pt(18)
        elif level == 2:
            run.font.size = Pt(13)
        else:
            run.font.size = Pt(12)
            
        return p

    # Tekil görsel ekleme fonksiyonu (Admin paneli yatay resimler için ideal)
    def add_single_img(img_name, caption, width_inch=4.5):
        img_path = os.path.join(img_dir, img_name)
        if os.path.exists(img_path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run()
            run.add_picture(img_path, width=Inches(width_inch))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(8)
            run_cap = p_cap.add_run(caption)
            run_cap.font.name = 'Times New Roman'
            run_cap.font.size = Pt(10)
            run_cap.italic = True
            print(f"Tekil resim eklendi: {img_name}")
        else:
            print(f"UYARI: Resim bulunamadı: {img_name}")

    # İki görseli yan yana tablosuz/borderless tablo ile ekleme fonksiyonu (Dikey mobil ekranlar için)
    def add_side_by_side_imgs(img1_name, img2_name, caption, width_inch=2.3):
        table = doc.add_table(rows=1, cols=2)
        # Tablo kenarlıklarını görünmez yap
        tblPr = table._element.xpath('w:tblPr')
        if tblPr:
            borders = parse_xml(r'<w:tblBorders %s><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/></w:tblBorders>' % nsdecls('w'))
            tblPr[0].append(borders)
        
        # 1. Hücre
        cell1 = table.cell(0, 0)
        p1 = cell1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_path1 = os.path.join(img_dir, img1_name)
        if os.path.exists(img_path1):
            p1.add_run().add_picture(img_path1, width=Inches(width_inch))
            
        # 2. Hücre
        cell2 = table.cell(0, 1)
        p2 = cell2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_path2 = os.path.join(img_dir, img2_name)
        if os.path.exists(img_path2):
            p2.add_run().add_picture(img_path2, width=Inches(width_inch))
            
        # Altyazı
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(10)
        run_cap = p_cap.add_run(caption)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(10)
        run_cap.italic = True
        print(f"Yan yana resimler eklendi: {img1_name} & {img2_name}")

    print("İçerik yazılıyor...")

    # --- ÖZET ---
    add_heading("ÖZET", level=1)
    add_p(
        "Bu çalışmada, perakende spor giyim sektöründe faaliyet gösteren Sporthink markasının müşteri sadakatini artırmak ve "
        "kullanıcı etkileşimini dijital kanallar aracılığıyla sürdürülebilir kılmak amacıyla geliştirilen bütünleşik bir mobil sağlık "
        "ve aktivite oyunlaştırma sistemi sunulmaktadır. Geliştirilen sistem, kullanıcıların günlük fiziksel aktivitelerini (adım sayılarını) "
        "mobil uygulama aracılığıyla takip ederek bunları oyun içi puanlara (Sporthink Puan) dönüştürmekte ve bu puanların mağaza içi "
        "kuponlar ve çeşitli avatar ödülleri ile takas edilmesini sağlamaktadır. Proje kapsamında, verilerin güvenliğini ve doğruluğunu "
        "sağlamak amacıyla gelişmiş bir adım doğrulama (anti-cheat) algoritması tasarlanmış ve veri suistimallerinin önüne geçilmiştir. "
        "Ayrıca, sistemin esnek ve sürdürülebilir bir yapıda kalması için tüm dinamiklerin (sandık düşme olasılıkları, görev ödülleri, günlük limitler vb.) "
        "tek bir merkezden yönetilebildiği kapsamlı bir web tabanlı admin yönetim paneli geliştirilmiştir. Mobil uygulama katmanında Flutter, "
        "servis ve kontrol katmanında Node.js/Express, veritabanı katmanında ise MySQL kullanılarak yüksek performanslı ve güvenli bir mimari "
        "elde edilmiştir. Bu çalışma, hem fiziksel mağaza trafiğini hem de dijital etkileşimi birleştiren hibrit bir müşteri ilişkileri "
        "yönetimi (CRM) modeli sunması açısından Yönetim Bilişim Sistemleri (YBS) literatürüne ve sektöre önemli katkılar sağlamaktadır."
    )
    add_p("Anahtar Kelimeler: Oyunlaştırma, Perakende Sadakat Sistemleri, Mobil Sağlık, Flutter, Anti-Cheat, Yönetim Bilişim Sistemleri.", italic=True)

    # --- GİRİŞ ---
    doc.add_page_break()
    add_heading("GİRİŞ", level=1)
    add_p(
        "Dijitalleşen dünyada, geleneksel perakende sektörü ciddi bir dönüşüm süreci geçirmektedir. Spor giyim ve ekipmanları "
        "sektörünün lider markalarından biri olan Sporthink, müşterileri ile sadece satış odaklı değil, aynı zamanda sağlıklı "
        "yaşam bilincini teşvik eden ve süreklilik arz eden bir bağ kurmayı hedeflemektedir. Mobil teknolojilerin ve giyilebilir "
        "sağlık cihazlarının yaygınlaşması, bireylerin günlük fiziksel aktivitelerini anlık olarak takip edebilmelerine olanak tanımaktadır. "
        "Bu bağlamda oyunlaştırma (gamification), kullanıcıların motivasyonunu artırarak belirli davranış kalıplarını (aktif kalma, spor yapma) "
        "düzenli hale getirmeleri için en etkili yöntemlerden biri olarak karşımıza çıkmaktadır."
    )
    add_p(
        "Geleneksel perakende sadakat sistemleri, genellikle yalnızca harcama yapıldığında puan kazandıran, statik ve kullanıcıda heyecan "
        "uyandırmayan pasif yapılardan oluşmaktadır. Müşteriler, alışveriş yapmadıkları dönemlerde markanın dijital uygulamalarını "
        "ziyaret etmemekte, bu da müşteri yaşam boyu değerinin (LTV) ve günlük aktif kullanıcı (DAU) oranlarının düşük kalmasına neden "
        "olmaktadır. Sporthink için tasarlanan bu proje, bu problemi kökten çözmek amacıyla fiziksel aktiviteyi doğrudan ödüllendiren, "
        "kullanıcılara kişiselleştirilebilir avatarlar ve sandık (loot chest) açma gibi oyun mekanikleri sunan, dinamik bir sadakat ekosistemi "
        "oluşturmayı amaçlamaktadır."
    )
    add_p(
        "Projenin hayata geçirilmesinde karşılaşılan en kritik zorluklardan biri, fiziksel adım verilerinin manipülasyona (fake step) "
        "açık olmasıdır. Kullanıcıların sahte yazılımlar veya cihaz sallama gibi yöntemlerle haksız puan elde etmesini engellemek amacıyla, "
        "sunucu tarafında çalışan özgün bir anti-cheat mekanizması kurgulanmıştır. Bu mekanizma, adımların frekansını, zaman damgalarını ve "
        "senkronizasyon sıklığını analiz ederek şüpheli aktiviteleri otomatik olarak askıya almaktadır."
    )
    add_p(
        "Ayrıca, projenin en büyük felsefesi esneklik ve sürdürülebilirliktir. Bir oyunlaştırma projesinin başarısı, kullanıcıların ilgisine "
        "göre kuralların sürekli güncellenmesini gerektirir. Bu doğrultuda, pazarlama veya BT departmanı yetkililerinin kod yazmasına gerek "
        "kalmadan, sistemdeki tüm parametreleri değiştirebileceği, sandıklardan çıkacak ödüllerin olasılık yüzdelerini (yüzde toplamını %100'e "
        "eşitleyen yardımcı araçlarla birlikte) güncelleyebileceği, global bildirimler gönderebileceği ve kullanıcı loglarını anlık izleyebileceği "
        "gelişmiş bir web tabanlı yönetim paneli kurgulanmıştır. Bu rapor kapsamında, sistemin tasarımı, analiz aşamaları, Sistem Geliştirme Yaşam Döngüsü "
        "(SGYD) adımları, teknik mimarisi ve elde edilen uygulama bulguları detaylandırılmıştır."
    )

    # --- BÖLÜM 1 ---
    doc.add_page_break()
    add_heading("BÖLÜM 1: MEVCUT SİSTEMİN TANIMI VE İNCELENMESİ", level=1)
    
    add_heading("1. Proje Konusu", level=2)
    add_p(
        "Bu projenin konusu, spor perakendeciliği sektöründe faaliyet gösteren Sporthink markası için adım tabanlı mobil sağlık verileri ile "
        "bütünleşik çalışan, anti-cheat (hile önleme) korumalı, kişiselleştirilebilir avatar ve oyun mağazası barındıran, arka plan operasyonlarının "
        "tamamen yönetici kontrol paneli üzerinden yapılandırıldığı dinamik bir sadakat ve oyunlaştırma platformunun tasarımı ve geliştirilmesidir."
    )

    add_heading("2. Amaç", level=2)
    add_p(
        "Projenin temel amacı, Sporthink müşterilerinin markayla olan ilişkisini sürekli kılmak, günlük aktif uygulama kullanımı alışkanlığı "
        "kazandırmak ve fiziksel aktiviteyi maddi ve manevi (dijital itibar/avatar) ödüllerle destekleyerek sağlıklı yaşamı teşvik etmektir. "
        "Bunun yanı sıra, yönetici tarafında kod bağımlılığını sıfıra indirerek oyun içi ekonominin dengelerini anlık olarak optimize "
        "edebilmeyi sağlayacak merkezi bir kontrol arayüzü sunmak hedeflenmiştir."
    )

    add_heading("3. Literatür Taraması", level=2)
    add_p(
        "Oyunlaştırma, oyun dışı ortamlarda oyun tasarım unsurlarının ve mekaniklerinin kullanılması olarak tanımlanmaktadır (Deterding vd., 2011). "
        "Özellikle Yu-kai Chou tarafından geliştirilen Octalysis Çerçevesi (Octalysis Framework), oyunlaştırmanın insan psikolojisi üzerindeki sekiz "
        "temel dürtüsünü (anlam, başarı, güçlendirme, sahiplik, sosyal etki, kıtlık, merak ve kayıp kaçınma) ortaya koymaktadır. Bu projede, "
        "kullanıcılara sunulan avatar özelleştirme 'sahiplik ve kendini ifade etme' dürtüsüne, liderlik tablosu 'sosyal etki ve rekabet' dürtüsüne, "
        "sandık açma sistemi ise 'merak ve öngörülemezlik' dürtüsüne hitap etmektedir."
    )
    add_p(
        "Literatürde yer alan mobil sağlık (mHealth) araştırmaları, bireysel aktivite takibinin oyun öğeleriyle birleştirildiğinde "
        "kullanıcı katılımını ve egzersiz yapma sıklığını ciddi oranda artırdığını göstermektedir (Hamari vd., 2014). Ancak, perakende sadakat programları "
        "ile mHealth uygulamalarının kesişimini inceleyen çalışmaların sınırlı olduğu görülmüştür. Bu proje, kullanıcıların adımlarını doğrudan "
        "fiziksel mağaza indirim kuponlarına dönüştürerek bu iki disiplini başarıyla entegre etmektedir."
    )

    add_heading("4. Problem Tanımı", level=2)
    add_p(
        "Mevcut perakende uygulamalarında temel problem, kullanıcıların yalnızca satın alım esnasında uygulamayla etkileşime girmesidir. "
        "Bu durum, yüksek müşteri kazanım maliyetlerine (CAC) karşın düşük elde tutma oranlarına (Retention Rate) yol açmaktadır. Ayrıca, mevcut "
        "adım sayar ödül uygulamalarında hileli adımların sisteme kolayca entegre edilmesi, şirketler için ciddi bir finansal risk ve maliyet "
        "oluşturmaktadır. Puan sisteminin dengesiz dağılımı ve yöneticilerin bu sistemleri dinamik olarak müdahale edip yönetememesi de "
        "operasyonel bir tıkanıklık yaratmaktadır."
    )

    add_heading("5. Araştırma Sorusu", level=2)
    add_p(
        "Araştırma Sorusu: 'Spor perakendeciliğinde müşteri sadakatini artırmak için tasarlanan adım tabanlı bir oyunlaştırma sisteminde, "
        "veri güvenliğini (anti-cheat) ve yönetsel esnekliği (admin paneli yönetimi) bir arada sağlayan bütünleşik bir sistem "
        "mimarisi nasıl kurgulanmalıdır?'"
    )

    add_heading("6. Proje Çerçevesi", level=2)
    add_p(
        "Bu proje, Yönetim Bilişim Sistemleri (YBS) prensiplerine uygun olarak üç ana bölüm halinde kurgulanmıştır. "
        "Bölüm 1'de projenin kuramsal temelleri, amaç ve sınırları çizilmiştir. Bölüm 2'de sistemin geliştirilmesinde kullanılan yazılım "
        "teknolojileri, mimari yapı ve Sistem Geliştirme Yaşam Döngüsü (SGYD) aşamaları açıklanmıştır. Bölüm 3'te ise geliştirilen sistemin "
        "mobil ve web arayüzlerine ait ekran görüntüleri eşliğinde elde edilen bulgular, entegrasyonlar ve detaylı sistem fonksiyonları "
        "sunulmuş, sonuç kısmında ise projenin akademik ve sektörel katkıları tartışılmıştır."
    )

    # --- BÖLÜM 2 ---
    doc.add_page_break()
    add_heading("BÖLÜM 2: YÖNTEM-METOD VE SİSTEM GELİŞTİRME YAŞAM DÖNGÜSÜ", level=1)
    
    add_p(
        "Projenin geliştirilmesinde modern yazılım mühendisliği yaklaşımları ve sistem analiz yöntemleri kullanılmıştır. "
        "Platformun mobil istemci tarafı Google tarafından geliştirilen, tek kod tabanıyla yüksek performanslı yerel arayüzler sunan "
        "Flutter SDK ve Dart dili ile kodlanmıştır. Sunucu (API) tarafı ise asenkron olay tabanlı mimarisi ve yüksek performansı sebebiyle "
        "Node.js runtime ortamında Express.js çatısı ile kurulmuştur. İlişkisel verilerin tutarlılığı, işlem güvenliği ve performanslı sorgulama "
        "ihtiyaçları için ilişkisel veritabanı yönetim sistemi (RDBMS) olarak MySQL tercih edilmiştir."
    )

    add_heading("Sistem Geliştirme Yaşam Döngüsünün (SGYD) 7 Aşaması", level=2)
    
    add_p("Proje, SGYD (SDLC) metodolojisine tam uyumlu olarak 7 temel aşamada analiz edilmiş ve hayata geçirilmiştir:", bold=True)

    add_p(
        "1. Problemlerin, Fırsatların ve Amaçların Tanımlanması: "
        "İlk aşamada Sporthink yetkilileriyle görüşmeler gerçekleştirilerek e-ticaret ve perakende mağazacılıkta müşteri kaybı ve düşük aktif "
        "kullanıcı sayısı gibi problemler tespit edilmiştir. Adım takibi entegrasyonunun, markanın sporcu ve aktif yaşam kimliğiyle uyuşan "
        "büyük bir pazarlama fırsatı sunduğu görülmüştür. Sistemin temel amacı; güvenli, yönetilebilir ve kullanıcı dostu bir oyunlaştırma "
        "deneyimi sunmak olarak belirlenmiş ve projenin teknik ve finansal fizibilite raporu hazırlanmıştır.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY
    )
    
    add_p(
        "2. Bilgi Gereksinimlerinin Belirlenmesi: "
        "Kullanıcıların mobil cihazlarından (iOS ve Android) sağlık verilerinin nasıl çekileceği araştırılmıştır. Google Health Connect ve Apple "
        "HealthKit entegrasyonu için gerekli veri şemaları analiz edilmiştir. Kullanıcılardan hangi izinlerin alınması gerektiği, veritabanında "
        "tutulacak kullanıcı profili, adım geçmişi, puan hareketleri ve avatar parçalarının veri türleri ve senkronizasyon zaman damgası "
        "gereksinimleri detaylı olarak listelenmiştir.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY
    )

    add_p(
        "3. Sistem İhtiyaçlarının Analizi: "
        "Toplanan bilgiler doğrultusunda sistemin mantıksal tasarımı yapılmıştır. Veri Akış Şemaları (DFD) çizilmiş ve sistemin sınırları "
        "belirlenmiştir. Anti-cheat koruması için adımların günlük limitleri, senkronizasyon başına maksimum adımlar ve şüpheli aktivite "
        "kriterleri matematiksel formüllere dökülmüştür. Yönetim panelinin hangi veri tablolarını (ürünler, sandıklar, kullanıcılar, denetim logları) "
        "CRUD işlemlerine tabi tutacağı belirlenmiştir.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY
    )

    add_p(
        "4. Önerilen Sistemin Tasarımı: "
        "Sistemin veri tabanı mimarisi (ERD şeması) tasarlanmıştır. Mobil uygulamanın modern, Sporthink kurumsal kimliğine uygun (koyu kırmızı "
        "ve siyah ağırlıklı premium karanlık mod) UI/UX tel kafes (wireframe) tasarımları oluşturulmuştur. Admin panelinin esnek yapısı için "
        "Tailwind CSS temelli duyarlı (responsive) şablonlar tasarlanmış, veritabanı tabloları arasındaki ilişkiler kurulmuştur.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY
    )

    add_p(
        "5. Yazılımın Geliştirilmesi ve Belge Yönetimi: "
        "Kodlama aşamasına geçilmiştir. Flutter tarafında modüler ekran yapıları, durum yönetimi (State Management), servis ve konfigürasyon "
        "katmanları oluşturulmuştur. Backend tarafında Express router mimarisi kurularak JWT tabanlı kimlik doğrulama, anti-cheat kontrolörleri, "
        "mağaza ve sandık algoritmaları kodlanmıştır. Kod güvenliği için tüm kod blokları git versiyon kontrol sistemi ile izlenmiş ve API "
        "dokümantasyonu oluşturulmuştur.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY
    )

    add_p(
        "6. Sistemin Test Edilmesi ve Sürdürülmesi: "
        "Sistem canlının önüne çıkmadan önce kapsamlı testlere tabi tutulmuştur. Sahte adım yazılımları ile anti-cheat mekanizmasının direnci "
        "test edilmiştir. Admin panelinde sandık ödül oranlarının dinamik değişiminin veritabanına ve oradan mobil uygulamaya doğru yansıdığı "
        "doğrulanmıştır. Çoklu cihaz testleri yapılarak ağ ve lokal IP değişikliklerinin senkronizasyonu sürdürülmüştür.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY
    )

    add_p(
        "7. Sistemin Gerçekleştirilmesi ve Değerlendirilmesi: "
        "Sistem sunucu ortamına kurulmuş, veritabanı göçleri (migrations) tamamlanmış ve mobil uygulamanın beta dağıtımı gerçekleştirilmiştir. "
        "Sistemin performansı, kullanıcıların günlük adımlarının hatasız senkronize olması ve yöneticilerin paneli kolaylıkla kullanabilmesi "
        "açısından değerlendirilmiş ve hedeflenen 8+ sayfalık kapsamlı teknik değerlendirme raporu ile proje başarıyla sonuçlandırılmıştır.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY
    )

    # --- BÖLÜM 3 ---
    doc.add_page_break()
    add_heading("BÖLÜM 3: BULGULAR VE UYGULAMA", level=1)
    
    add_p(
        "Geliştirilen Sporthink Aktivite ve Oyunlaştırma Platformu, iki ana arayüzden oluşmaktadır: Müşterilerin kullandığı ve "
        "fiziksel adımlarını puana dönüştürdükleri 'Sporthink Mobil Uygulaması' ve marka yöneticilerinin tüm sistemi kod bağımsız yönettikleri "
        "'Sporthink Web Yönetim Paneli'. Bu bölümde, uygulamanın ekran çıktıları ve işlevsel bulguları detaylı şekilde incelenmiştir."
    )

    add_heading("1. Mobil Uygulama Modülleri ve Ekran Çıktıları", level=2)
    
    add_p("Mobil uygulama dikey ekran tasarımları yan yana gruplandırılarak görsel şölen ve sayfa verimliliği bir arada sunulmuştur:", bold=True)

    add_p("A. Giriş Ekranı ve Ana Sayfa Aktivite Takibi:", bold=True)
    add_p(
        "Kullanıcıların sisteme güvenli bir şekilde dahil olabilmesi için JWT tabanlı kimlik doğrulama akışı kurulmuştur. "
        "Ana sayfa, kullanıcının o gün attığı adım sayısını, bu hafta attığı toplam adım sayısını ve mevcut aktif serisini (streak) "
        "gösteren dairesel ilerleme grafiklerine ev sahipliği yapmaktadır."
    )
    add_side_by_side_imgs("app-giriş.jpeg", "app-anasayfa1.jpeg", "Şekil 1: Mobil Uygulama Giriş Ekranı ve Ana Sayfa Genel Görünümü")

    add_p("B. Dinamik Bilgi Akışı ve Avatar Özelleştirme Modülü:", bold=True)
    add_p(
        "Health Connect entegrasyonu ile senkronize edilen adımlar anlık olarak sunucuya aktarılır. Kullanıcılar kazandıkları puanlar "
        "ve açtıkları sandıklardan kazandıkları parçalar ile avatarlarını tamamen kişiselleştirebilmektedir."
    )
    add_side_by_side_imgs("app-anasyafa2.jpeg", "app-avatar.jpeg", "Şekil 2: Sağlık Bilgi Akışı / İpuçları ve Avatar Düzenleyici Arayüzleri")

    add_p("C. Puan Mağazası ve Günlük Görevler:", bold=True)
    add_p(
        "Puan mağazasında indirim kuponları ve hediye sandıklar listelenmektedir. Günlük görevler ise kullanıcılara her gün tamamlamaları "
        "gereken yeni hedefler (örneğin 10.000 adım) atayarak onları aktif tutmaktadır."
    )
    add_side_by_side_imgs("app-dükkan.jpeg", "app-görevler.jpeg", "Şekil 3: Puan Mağazası (Sandık Satın Alımı) ve Günlük Görev Takip Paneli")

    add_p("D. Liderlik Tablosu ve Sosyal Akış:", bold=True)
    add_p(
        "Haftalık adım sıralamasında üst sıralara tırmanan kullanıcılar rekabetin tadını çıkarırken, Sosyal Akış ekranında takip ettikleri "
        "arkadaşlarının aktivite ve rozet durumlarını görebilmektedir."
    )
    add_side_by_side_imgs("app-sıralama.jpeg", "app-feed.jpeg", "Şekil 4: Rekabetçi Liderlik Tablosu ve Sosyal Paylaşım Akışı")

    add_p("E. Kullanıcı Profili ve Bildirim Merkezi:", bold=True)
    add_p(
        "Kullanıcı profilinde kazanılan rozetler ve seviye detayları yer almaktadır. Bildirim merkezinde ise yöneticilerin panel üzerinden "
        "gönderdiği kampanya duyuruları ve sistem uyarıları anlık olarak görüntülenir."
    )
    add_side_by_side_imgs("app-profil.jpeg", "app-bildirim.jpeg", "Şekil 5: Kullanıcı Başarı Profili ve Bildirim Geçmişi Ekranları")

    add_heading("2. Web Yönetim Paneli Modülleri ve Ekran Çıktıları", level=2)
    
    add_p("Yöneticilerin sistemi kontrol ettiği web paneli yatay ekran çıktıları en-boy oranı korunarak tekil olarak yerleştirilmiştir:", bold=True)

    add_p("A. Yönetim Merkezi (Dashboard):", bold=True)
    add_p(
        "Yönetim paneli ana sayfası olan Dashboard; toplam kullanıcı sayısı, sistem genelindeki toplam puan akışı, aktif banlı kullanıcı sayısı "
        "ve sistem yöneticilerinin gerçekleştirdiği tüm operasyonların güvenlik denetim loglarını (audit logs) sunmaktadır."
    )
    add_single_img("admin-dashboard.jpeg", "Şekil 6: Web Yönetici Kontrol Paneli Ana Ekranı (Dashboard)", width_inch=4.8)

    add_p("B. Kullanıcı Yönetimi, Log Görüntüleme ve Güvenlik Aksiyonları:", bold=True)
    add_p(
        "Yöneticiler kullanıcı detaylarını görebilmekte ve şüpheli hesapları banlayabilmektedir. Eklenen yeni özellikle, 'Log' butonuyla "
        "kullanıcının adım geçmişi, detaylı puan hareketleri (Points Ledger) ve kazandığı rozetler AJAX ile anında listelenebilmektedir."
    )
    add_single_img("admin-kullanıcılar.jpeg", "Şekil 7: Kullanıcı Detaylı Log Arama ve Güvenlik Yönetim Arayüzü", width_inch=4.8)

    add_p("C. Sandık ve Ödül Olasılık Yapılandırması:", bold=True)
    add_p(
        "Sandıklardan hangi avatar parçasının ne kadar ihtimalle düşeceği tamamen dinamiktir. Akıllı olasılık hesaplama motoru kalan "
        "olasılığı otomatik göstererek toplamın tam %100 olmasını garanti eder ve sistemin dengesini korur."
    )
    add_single_img("admin-ayarlar.jpeg", "Şekil 8: Sandık Drop Olasılıkları ve Sistem Ayarları Arayüzü", width_inch=4.8)

    # --- SONUÇ VE ÖNERİLER ---
    doc.add_page_break()
    add_heading("SONUÇ VE ÖNERİLER", level=1)
    add_p(
        "Bu proje kapsamında, perakende sektörünün öncülerinden Sporthink için bütünsel, güvenli ve yönetici dostu bir mobil sağlık "
        "ve aktivite oyunlaştırma platformu başarıyla hayata geçirilmiştir. Tasarlanan anti-cheat adım doğrulama algoritmaları sayesinde "
        "sistemin suistimal edilmesi engellenmiş ve adil bir rekabet ortamı tesis edilmiştir. Geliştirilen web tabanlı admin yönetim paneli "
        "sayesinde, pazarlama yetkilileri oyun içi tüm parametrelere, ödül olasılıklarına ve kullanıcı loglarına doğrudan müdahale edebilir hale "
        "getirilmiştir. Bu durum, BT departmanının üzerindeki operasyonel yükü sıfırlayarak işletmeye muazzam bir çeviklik kazandırmıştır."
    )
    
    add_heading("Yönetim Bilişim Sistemleri (YBS) Alanına Katkısı", level=2)
    add_p(
        "Yönetim Bilişim Sistemleri, teknoloji, insan ve organizasyon süreçlerini birleştiren disiplinler arası bir alandır. "
        "Bu proje, YBS disiplininin en temel felsefesini temsil etmektedir. Müşteri İlişkileri Yönetimi (CRM) ve e-ticaret sadakat "
        "süreçleri (Organizasyon), mobil adım sayar ve güvenli API entegrasyonlarıyla (Teknoloji) birleştirilerek, kullanıcıların "
        "sağlıklı yaşam motivasyonu (İnsan) aracılığıyla işletmeye değer katması sağlanmıştır."
    )
    add_p(
        "Geliştirilen sistem, yönetsel düzeyde bir Karar Destek Sistemi (DSS) ve Yönetim Raporlama Sistemi (MIS) olarak çalışmaktadır. "
        "Yöneticiler, kullanıcı loglarını ve genel istatistikleri inceleyerek hangi kampanya katsayılarının daha etkili olduğunu analiz "
        "edebilmekte ve oyun içi ekonomi kararlarını bu verilere dayanarak rasyonel şekilde alabilmektedir. Sonuç olarak, bu çalışma "
        "bilgi teknolojilerinin bir işletmenin pazarlama stratejilerine nasıl doğrudan kaldıraç etkisi yaratabileceğinin somut bir kanıtıdır."
    )

    add_heading("Gelecek Çalışmalar ve Öneriler", level=2)
    add_p(
        "Sistemin gelecekte daha ileri bir seviyeye taşınması için şu geliştirmeler önerilmektedir: "
        "1) Kullanıcıların kendi aralarında gruplar kurarak grup bazlı adım yarışmaları düzenleyebilmesi (Sosyal Etkileşimin artırılması), "
        "2) Yapay zeka destekli bir öneri motoru entegre edilerek, kullanıcının geçmiş adım verilerine göre kişiselleştirilmiş günlük egzersiz "
        "ve ürün tavsiyelerinde bulunulması, "
        "3) Akıllı saatler (Apple Watch, Garmin, Fitbit vb.) için yerel uygulamalar geliştirilerek entegrasyon derinliğinin artırılması."
    )

    # --- REFERANSLAR ---
    doc.add_page_break()
    add_heading("REFERANSLAR", level=1)
    add_p("Chou, Y. K. (2015). Actionable gamification: Beyond points, badges, and leaderboards. Octalysis Media.", align=WD_ALIGN_PARAGRAPH.LEFT)
    add_p("Deterding, S., Dixon, D., Khaled, R., & Nacke, L. (2011). From game design elements to gamefulness: defining gamification. In Proceedings of the 15th international academic MindTrek conference: Envisioning future media environments (pp. 9-15).", align=WD_ALIGN_PARAGRAPH.LEFT)
    add_p("Hamari, J., Koivisto, J., & Sarsa, H. (2014). Does gamification work?--a literature review of empirical studies on gamification. In 2014 47th Hawaii international conference on system sciences (pp. 3025-3034). IEEE.", align=WD_ALIGN_PARAGRAPH.LEFT)
    add_p("Werbach, K., & Hunter, D. (2012). For the win: How game thinking can revolutionize your business. Wharton Digital Press.", align=WD_ALIGN_PARAGRAPH.LEFT)

    print("Kaydediliyor...")
    doc.save(output_path)
    print(f"BAŞARILI: Rapor optimize edildi ve kaydedildi: {output_path}")

if __name__ == '__main__':
    create_report()
